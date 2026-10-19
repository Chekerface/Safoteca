#!/usr/bin/env python3
"""Completa el campo «final_feliz» (sí / no) de las obras.

Reglas automáticas (solo sobre obras que aún no tienen dato):
  * Tiene el aviso «Síndrome de la lesbiana muerta»  ->  final no feliz.

Lo demás no se puede deducir, así que se saca a un Excel para revisarlo a mano y luego se vuelve a cargar.

Uso:
    python tools/completar_final_feliz.py --simular                  # solo informa
    python tools/completar_final_feliz.py                            # aplica las reglas
    python tools/completar_final_feliz.py --pendientes pendientes.xlsx   # además crea el Excel de pendientes
    python tools/completar_final_feliz.py --aplicar revisado.xlsx    # carga la columna «¿Final feliz?» (Sí / No)
Después: python tools/exportar_json.py
"""
import argparse, shutil, sqlite3
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
REGLAS_NO = ["Síndrome de la lesbiana muerta"]
# Pistas que se enseñan en el Excel para ayudar a revisar (no se usan para decidir nada)
PISTAS_AVISOS = ["Muerte de personaje importante", "Muerte", "Suicidio", "Desamor", "Cuernos", "Infidelidad", "Cancelada"]
PISTAS_TEMATICAS = ["Heartwarming", "No heartwarming", "Romance", "Comedia romántica", "Tragedia", "Drama"]


def copia(db_ruta):
    destino = RAIZ / "data" / "privado" / "copias" / f"safoteca_antes_final_feliz_{datetime.now():%Y%m%d_%H%M%S}.db"
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(db_ruta, destino)
    print("Copia de seguridad:", destino.relative_to(RAIZ))


def titulo(f):
    return f[1] or f[2]


def reglas(db, guardar):
    hechas = []
    for aviso in REGLAS_NO:
        filas = db.execute("""SELECT o.id, o.titulo_es, o.titulo_original, o.final_feliz FROM obra o
            JOIN obra_aviso x ON x.obra_id = o.id JOIN aviso a ON a.id = x.aviso_id
            WHERE a.nombre = ? AND o.tipo IN ('serie','pelicula')""", (aviso,)).fetchall()
        for f in filas:
            if f[3] is None:
                if guardar:
                    db.execute("UPDATE obra SET final_feliz = 0 WHERE id = ?", (f[0],))
                hechas.append((f[0], titulo(f), f"«{aviso}» -> final no feliz"))
            elif f[3] == 1:
                hechas.append((f[0], titulo(f), f"OJO: tiene «{aviso}» pero consta final feliz (no se toca)"))
    return hechas


def pendientes(db, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Pendientes"
    ws.append(["id", "Título", "Tipo", "Año", "Estado", "Pistas en avisos", "Pistas en temáticas", "Sinopsis", "¿Final feliz? (Sí / No)"])
    filas = db.execute("""SELECT id, titulo_es, titulo_original, tipo, anio, estado, sinopsis FROM obra
        WHERE tipo IN ('serie','pelicula') AND final_feliz IS NULL
        ORDER BY (estado = 'publicada') DESC, COALESCE(titulo_es, titulo_original) COLLATE NOCASE""").fetchall()
    for f in filas:
        av = [r[0] for r in db.execute("SELECT a.nombre FROM obra_aviso x JOIN aviso a ON a.id = x.aviso_id WHERE x.obra_id = ?", (f[0],))]
        te = [r[0] for r in db.execute("SELECT t.nombre FROM obra_tematica x JOIN tematica t ON t.id = x.tematica_id WHERE x.obra_id = ?", (f[0],))]
        ws.append([f[0], titulo(f), "Serie" if f[3] == "serie" else "Película", f[4], f[5],
                   ", ".join(a for a in PISTAS_AVISOS if a in av), ", ".join(t for t in PISTAS_TEMATICAS if t in te), f[6] or "", ""])
    rel = PatternFill("solid", fgColor="4B2E83")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = rel; c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, w in zip("ABCDEFGHI", (7, 36, 10, 7, 11, 28, 28, 90, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)
    return len(filas)


def aplicar(db, ruta):
    from openpyxl import load_workbook
    ws = load_workbook(ruta).worksheets[0]
    cab = [str(c.value or "").strip().lower() for c in ws[1]]
    i_id = cab.index("id"); i_dec = next(i for i, c in enumerate(cab) if c.startswith("¿final feliz"))
    si, no, saltadas = 0, 0, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        v = (str(r[i_dec]).strip().lower() if r[i_dec] is not None else "")
        if not v:
            continue
        if v in ("sí", "si", "s", "1", "true"):
            db.execute("UPDATE obra SET final_feliz = 1 WHERE id = ?", (int(r[i_id]),)); si += 1
        elif v in ("no", "n", "0", "false"):
            db.execute("UPDATE obra SET final_feliz = 0 WHERE id = ?", (int(r[i_id]),)); no += 1
        else:
            saltadas.append((r[i_id], r[i_dec]))
    print(f"Cargadas: {si} con final feliz y {no} sin final feliz.")
    for s in saltadas:
        print("  Valor no reconocido (se salta):", s)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=str(RAIZ / "safoteca.db"))
    ap.add_argument("--simular", action="store_true")
    ap.add_argument("--pendientes", metavar="EXCEL")
    ap.add_argument("--aplicar", metavar="EXCEL")
    a = ap.parse_args()
    db = sqlite3.connect(a.db)
    if a.aplicar:
        if not a.simular:
            copia(a.db)
        aplicar(db, a.aplicar)
        if a.simular:
            db.rollback(); print("(Simulación: no se ha guardado nada.)")
        else:
            db.commit()
        return
    if not a.simular:
        copia(a.db)
    hechas = reglas(db, guardar=not a.simular)
    for h in hechas:
        print(f"  {h[1]}: {h[2]}")
    print(f"{len(hechas)} obras tocadas por las reglas." + (" (Simulación: no se ha guardado nada.)" if a.simular else ""))
    if not a.simular:
        db.commit()
    if a.pendientes:
        n = pendientes(db, a.pendientes)
        print(f"Excel de pendientes con {n} obras:", a.pendientes)


if __name__ == "__main__":
    main()
