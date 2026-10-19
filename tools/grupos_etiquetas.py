#!/usr/bin/env python3
"""Grupos de temáticas y avisos: un subgrupo (Racismo) pertenece a un grupo (Discriminación).

Las obras que tienen un subgrupo reciben solas su grupo, y el grupo del grupo (Racismo -> Discriminación -> Violencia).
Esto lo hace la propia base de datos con disparadores (db/schema.sql), así que también vale para todo lo que se añada
después: formulario, importaciones de Excel o cambios a mano.

La lista de grupos vive en dos CSV (columnas hijo,grupo):
    data/mapeos/grupos_tematicas.csv   data/mapeos/grupos_avisos.csv

Uso:
    python tools/grupos_etiquetas.py --excel Listado.xlsx   # crea/actualiza los CSV desde el Excel de revisión
    python tools/grupos_etiquetas.py --simular              # solo informa
    python tools/grupos_etiquetas.py                        # carga los CSV en la base de datos (hace copia antes)
Después: python tools/exportar_json.py y python tools/generar_formulario.py
Es seguro repetirlo.
"""
import argparse, csv, shutil, sqlite3, sys
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
MAPEOS = RAIZ / "data" / "mapeos"
HOJAS = {"tematica": ("Temáticas", "grupos_tematicas.csv"), "aviso": ("Avisos (TW)", "grupos_avisos.csv")}
# Erratas del Excel y nombres que cambian al aplicar las decisiones de unificar_etiquetas.py
CORRECCIONES = {"Driscriminación": "Discriminación", "Padre lgbtfobo": "Padre lgbtfóbico", "Padre lgtbfobo": "Padre lgbtfóbico",
                "Mentiras y Engaño": "Engaños y Mentiras"}


def leer_excel(ruta):
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True)
    for tabla, (hoja, csvname) in HOJAS.items():
        ws = wb[hoja]
        pares = []
        for r in list(ws.iter_rows(values_only=True))[1:]:
            grupo, sub = (r[0] or "").strip(), (r[1] or "").strip()
            if not grupo or not sub:
                continue
            grupo, sub = CORRECCIONES.get(grupo, grupo), CORRECCIONES.get(sub, sub)
            if grupo.lower() == sub.lower():
                print(f"  ({tabla}) «{sub}» tiene como grupo a sí misma: se ignora")
                continue
            pares.append((sub, grupo))
        with open(MAPEOS / csvname, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f); w.writerow(["hijo", "grupo"]); w.writerows(sorted(set(pares), key=lambda p: (p[1].lower(), p[0].lower())))
        print(f"{csvname}: {len(set(pares))} relaciones")


def leer_csv(nombre):
    ruta = MAPEOS / nombre
    if not ruta.exists():
        return []
    with open(ruta, encoding="utf-8", newline="") as f:
        return [(r["hijo"].strip(), r["grupo"].strip()) for r in csv.DictReader(f) if r.get("hijo") and r.get("grupo")]


def aplicar(db, informe):
    db.executescript((RAIZ / "db" / "schema.sql").read_text(encoding="utf-8"))  # crea tablas y disparadores si faltan
    for tabla, (_, csvname) in HOJAS.items():
        enlace, tg = "obra_" + tabla, tabla + "_grupo"
        et = "temática" if tabla == "tematica" else "aviso"
        ids = {}

        def ident(nombre, crear):
            r = db.execute(f"SELECT id FROM {tabla} WHERE nombre = ?", (nombre,)).fetchone()
            if r:
                return r[0]
            if not crear:
                return None
            extra = ", es_tw" if tabla == "aviso" else ""
            vals = ", 1" if tabla == "aviso" else ""
            db.execute(f"INSERT INTO {tabla}(nombre{extra}) VALUES (?{vals})", (nombre,))
            informe.append(f"{et}: se crea el grupo «{nombre}» en la lista")
            return db.execute(f"SELECT id FROM {tabla} WHERE nombre = ?", (nombre,)).fetchone()[0]

        pares = leer_csv(csvname)
        if not pares:
            continue
        nuevas = 0
        for hijo, grupo in pares:
            h = ident(hijo, False)
            if h is None:
                informe.append(f"{et}: «{hijo}» no existe en la base de datos (se salta su grupo «{grupo}»)")
                continue
            g = ident(grupo, True)
            if db.execute(f"INSERT OR IGNORE INTO {tg}(hijo_id, padre_id) VALUES (?, ?)", (h, g)).rowcount:
                nuevas += 1
        # por si un grupo ya existía y alguna obra aún no lo tiene
        db.execute(f"INSERT OR IGNORE INTO {enlace}(obra_id, {tabla}_id) SELECT x.obra_id, a.padre_id FROM {enlace} x JOIN {tabla}_ancestro a ON a.hijo_id = x.{tabla}_id")
        informe.append(f"{et}: {nuevas} relaciones nuevas de grupo ({len(pares)} en el CSV)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=str(RAIZ / "safoteca.db"))
    ap.add_argument("--excel", metavar="EXCEL")
    ap.add_argument("--simular", action="store_true")
    a = ap.parse_args()
    if a.excel:
        leer_excel(a.excel)
        return
    db = sqlite3.connect(a.db)
    db.execute("PRAGMA foreign_keys = ON")
    antes = {t: db.execute(f"SELECT COUNT(*) FROM obra_{t}").fetchone()[0] for t in ("tematica", "aviso")} if db.execute("SELECT name FROM sqlite_master WHERE name='obra_tematica'").fetchone() else {}
    if not a.simular:
        copia = RAIZ / "data" / "privado" / "copias" / f"safoteca_antes_grupos_{datetime.now():%Y%m%d_%H%M%S}.db"
        copia.parent.mkdir(parents=True, exist_ok=True); shutil.copy(a.db, copia)
        print("Copia de seguridad:", copia.relative_to(RAIZ))
    informe = []
    aplicar(db, informe)
    print("\n".join(informe))
    for t, n in antes.items():
        print(f"{t}: {n} asignaciones obra-etiqueta antes -> {db.execute(f'SELECT COUNT(*) FROM obra_{t}').fetchone()[0]} después")
    if a.simular:
        db.rollback(); print("\n(Simulación: no se ha guardado nada.)")
    else:
        db.commit(); print("\nListo. Ahora ejecuta tools/exportar_json.py y tools/generar_formulario.py")


if __name__ == "__main__":
    sys.exit(main())
